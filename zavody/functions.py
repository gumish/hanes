
import csv
from datetime import datetime
from io import TextIOWrapper

from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment

from kluby.models import Klub
from lide.models import Clenstvi, Clovek
from zavodnici.models import Zavodnik

from .models import Kategorie, Rocnik, Sport, Zavod
from .templatetags import custom_filters
from .templatetags.custom_filters import desetiny_sekundy, rozsah_narozeni


def _csv_reader(soubor):
    soubor_text = TextIOWrapper(soubor, encoding='windows-1250')
    csv_reader = csv.reader(soubor_text, delimiter=';')
    for row in csv_reader:
        yield [cell.strip() for cell in row]

def _pohlavi(slovo):
    "preklada muzi/zeny"
    if slovo:
        return slugify(slovo[0])
    else:
        return ''

# -------
# IMPORTY
# -------

def csv_kategorie_import(soubor):
    "pouziti pri zakladani noveho rocniku"
    chyby = []
    kategorie_list = []
    for i, radek in enumerate(_csv_reader(soubor), start=1):
        print(radek)
        if i >= 5:
            try:
                kategorie = Kategorie(
                    nazev=radek[0],
                    znacka=radek[1],
                    pohlavi=radek[2],
                    vek_od=int(radek[3]) if radek[3] else None,
                    vek_do=int(radek[4]) if radek[4] else None,
                    delka_trate=radek[5],
                    startovne=(radek[6] if radek[6] else 0) or 0,
                )
                kategorie_list.append(kategorie)
            except Exception as error:
                chyby.append(f'#{i} {error}')
    return (kategorie_list, chyby)

def rocnik_import(soubor):
    "import vysledku EXPORT VYSLEDKOVE LISTINY (fce exportuj_vysledky)"

    def _text_na_cas(text, i=None):
        if text:
            format_casu = '%H:%M:%S,%f' if ',' in text else '%H:%M:%S'
            return datetime.strptime(text, format_casu).time()
        else:
            return None

    zpravy = []
    predchozi_radek = None
    for i, radek in enumerate(_csv_reader(open(soubor)), start=1):
        if not radek:
            # prazdny radek
            predchozi_radek = 'prazdny'
        else:
            print('*'*10)
            if i == 1:
                # hlavicka
                sport, created = Sport.objects.get_or_create(nazev=radek[2])
                if created:
                    zpravy.append(f'#{i} uložen nový sport: {sport}')
                zavod, created = Zavod.objects.get_or_create(nazev=radek[0], sport=sport)
                if created:
                    zpravy.append(f'#{i} uložen nový závod: {zavod}')
                rocnik, created = Rocnik.objects.get_or_create(
                    zavod=zavod,
                    datum=datetime.strptime(radek[1], '%Y-%m-%d').date()
                )
            elif radek[0] == 'pořadí':
                # hlavicka zavodniku
                print(i, 'hlavicka zavodniku')
                predchozi_radek = 'hlavicka zavodniku'
            elif predchozi_radek == 'prazdny' and radek[0]:
                # csv_kategorie_import
                print(i, 'kategorie')
                od_do = list(map(int, radek[2].split(' - ')))
                od_do = [rocnik.datum.year - x for x in od_do]
                kategorie, created = Kategorie.objects.get_or_create(
                    znacka=radek[0],
                    nazev=radek[1],
                    pohlavi=_pohlavi(radek[3]),
                    vek_od=od_do[1],
                    vek_do=od_do[0],
                    delka_trate=radek[4],
                    spusteni_stopek=_text_na_cas(radek[5], i),
                    rocnik=rocnik
                )
                predchozi_radek = 'kategorie'
                if created:
                    zpravy.append(f'#{i} uložena nová kategorie: {kategorie}')
            elif predchozi_radek in ('hlavicka zavodniku', 'zavodnik'):
                # Zavodnik
                print(i, 'zavodnik')
                klub, created = Klub.objects.get_or_create(
                    slug=slugify(radek[5]),
                    defaults={'nazev': radek[5].strip()})
                if created:
                    zpravy.append(f'#{i} uložen nový klub: {klub}')
                clovek, created = Clovek.objects.get_or_create(
                    jmeno=radek[3],
                    prijmeni=radek[2],
                    narozen=int(radek[4])
                )
                if created:
                    zpravy.append(f'#{i} uložen nový člověk: {clovek}')
                try:
                    _clenstvi, _created = Clenstvi.objects.get_or_create(
                        clovek=clovek,
                        klub=klub
                    )
                except:
                    pass
                _zavodnik, created = Zavodnik.objects.get_or_create(
                    rocnik=rocnik,
                    clovek=clovek,
                    klub=klub,
                    kategorie_temp=kategorie,
                    cislo=radek[1],
                    startovni_cas=_text_na_cas(radek[6], i),
                    cilovy_cas=_text_na_cas(radek[7], i),
                    nedokoncil=radek[9],
                    odstartoval=radek[10]
                )
                predchozi_radek = 'zavodnik'
    return zpravy


# -------
# EXPORTY
# -------

def exportuj_startovku(response, rocnik, ordering_str):
    def _zavodnici_write(zavodnici):
        for zavodnik in zavodnici:
            writer.writerow([
                zavodnik.cislo,
                zavodnik.clovek.prijmeni,
                zavodnik.clovek.jmeno,
                zavodnik.clovek.narozen,
                zavodnik.klub.nazev if zavodnik.klub else '',
                desetiny_sekundy(zavodnik.startovni_cas),
                desetiny_sekundy(zavodnik.cilovy_cas),
                zavodnik.nedokoncil or desetiny_sekundy(zavodnik.vysledny_cas),
                zavodnik.odstartoval
            ])

    enctype = 'cp1250'
    writer = csv.writer(response, delimiter=';')
    kategorie_list = rocnik.kategorie_list(ordering_str)
    nezarazeni = rocnik.nezarazeni(ordering_str)

    writer.writerow([rocnik.__str__()])
    writer.writerow([])
    writer.writerow([])

    for kategorie, zavodnici in kategorie_list:
        writer.writerow([
            kategorie.znacka,
            kategorie.nazev,
            custom_filters.rozsah_narozeni(kategorie.rozsah_narozeni()),
            kategorie.get_pohlavi_display(),
            kategorie.delka_trate
        ])
        _zavodnici_write(zavodnici)
        writer.writerow([])
        writer.writerow([])

    if nezarazeni:
        writer.writerow([
            '',
            'nezařazení'
        ])
        _zavodnici_write(nezarazeni)

    return response


def exportuj_vysledky(response, rocnik):
    def _zavodnici_write(zavodnici):
        for zavodnik in zavodnici:
            writer.writerow([
                zavodnik.poradi_v_kategorii(),
                zavodnik.cislo,
                zavodnik.clovek.prijmeni,
                zavodnik.clovek.jmeno,
                zavodnik.clovek.narozen,
                zavodnik.klub.nazev if zavodnik.klub else '',
                desetiny_sekundy(zavodnik.startovni_cas),
                desetiny_sekundy(zavodnik.cilovy_cas),
                desetiny_sekundy(zavodnik.vysledny_cas),
                zavodnik.nedokoncil,
                zavodnik.odstartoval,
                zavodnik.poradi_na_trati()
            ])

    enctype = 'cp1250'
    writer = csv.writer(response, delimiter=';')
    kategorie_list = rocnik.kategorie_list()

    writer.writerow([
        rocnik.zavod.nazev,
        rocnik.datum,
        rocnik.zavod.sport.__str__()
    ])
    writer.writerow([])
    writer.writerow([])

    for kategorie, zavodnici in kategorie_list:
        writer.writerow([
            kategorie.znacka,
            kategorie.nazev,
            custom_filters.rozsah_narozeni(kategorie.rozsah_narozeni()),
            kategorie.get_pohlavi_display(),
            kategorie.delka_trate,
            kategorie.spusteni_stopek
        ])
        writer.writerow([
            'pořadí',
            'číslo',
            'příjmení',
            'jméno',
            'nar.',
            'klub',
            'startovní čas',
            'cílový čas',
            'výsledný čas',
            'nedokončil',
            'odstartoval',
            'na trati',
        ])
        _zavodnici_write(zavodnici)
        writer.writerow([])
        writer.writerow([])

    return response


def exportuj_kategorie(response, rocnik):
    enctype = 'cp1250'
    writer = csv.writer(response, delimiter=';')

    writer.writerow([
        'Název závodu',
        'Sport'
    ])
    writer.writerow([
        rocnik.__str__(),
        rocnik.zavod.sport.__str__()
    ])
    writer.writerow([])
    writer.writerow([
        'Název',
        'Značka',
        'Pohlaví',
        'Věk od včetně',
        'Věk do včetně',
        'Délka tratě',
        'Startovné'
    ])
    for kategorie in rocnik.kategorie.all():
        writer.writerow([
            kategorie.nazev,
            kategorie.znacka,
            kategorie.pohlavi,
            kategorie.vek_od,
            kategorie.vek_do,
            kategorie.delka_trate,
            kategorie.startovne or 0,
        ])

    return response


def exportuj_kategorie_xlsx(response, rocnik):
    wb = Workbook()
    wb.remove(wb.active)

    headers = [
        'Pořadí', 'Start. číslo', 'Kód závodníka', 'Příjmení', 'Jméno',
        'Rok nar.', 'Pohlaví', 'Stát', 'FIS kód', 'Název klubu', 'Čas'
    ]

    for kategorie in rocnik.kategorie.all():
        znacka = kategorie.znacka or str(kategorie.id)
        sheet_title = f'{znacka} - {kategorie.nazev}'
        ws = wb.create_sheet(title=sheet_title[:31])

        datum_str = f'{rocnik.datum.day}. {rocnik.datum.month}. {rocnik.datum.year}'
        rozsah = rozsah_narozeni(kategorie.rozsah_narozeni())
        titulek = f'{datum_str} \u2013 {rocnik.zavod.nazev}\n{kategorie.nazev} ({rozsah})'
        ws['A1'] = titulek
        ws.merge_cells('A1:K1')
        ws['A1'].alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 40

        for col, header in enumerate(headers, 1):
            ws.cell(row=2, column=col, value=header)

        zavodnici = kategorie.serazeni_zavodnici(razeni=None)
        for row_idx, zavodnik in enumerate(zavodnici, 3):
            ws.cell(row=row_idx, column=1, value=zavodnik.poradi_v_kategorii())
            ws.cell(row=row_idx, column=2, value=zavodnik.cislo)
            ws.cell(row=row_idx, column=4, value=zavodnik.clovek.prijmeni)
            ws.cell(row=row_idx, column=5, value=zavodnik.clovek.jmeno)
            ws.cell(row=row_idx, column=6, value=zavodnik.clovek.narozen)
            ws.cell(
                row=row_idx, column=7,
                value=zavodnik.clovek.get_pohlavi_display() if zavodnik.clovek.pohlavi else '')
            if zavodnik.clovek.stat:
                ws.cell(row=row_idx, column=8, value=zavodnik.clovek.stat.zkratka)
            ws.cell(row=row_idx, column=10, value=zavodnik.klub.nazev if zavodnik.klub else '')
            ws.cell(row=row_idx, column=11, value=desetiny_sekundy(zavodnik.vysledny_cas))

    wb.save(response)
    return response
